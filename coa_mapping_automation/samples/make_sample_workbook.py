#!/usr/bin/env python3
"""Generate samples/COA_Mapping_Template.xlsx - the layout Finance fills in."""
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "COA_Mapping_Template.xlsx")

wb = openpyxl.Workbook()
ins = wb.active
ins.title = "Instructions"
for line in [
    "COA mapping upload - drop the saved .xlsx into the loader inbox.",
    "Tab names do not matter; tabs are recognised by their header row.",
    "Target Values tab: one row per target value. MAPPING SET NAME + TARGET 1 identify the row.",
    "  Existing row -> other targets are updated; new row -> inserted. Unknown mapping set -> created.",
    "GL Mapping tab: one row per GL account. MAPPING SET NAME + SEGMENT 1 + SEGMENT 2 identify the row.",
    "  Every target used here must exist on the Target Values tab or already in the system.",
    "Blank cells keep the value already loaded. Format segment columns as Text to keep leading zeros.",
    "Any error rejects the whole file; the error report lists the Excel line of each problem.",
]:
    ins.append([line])

head = Font(bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="1F4E78")


def sheet(title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = head, fill
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
        col[0].number_format = "@"
    return ws


sheet("Target Values",
      ["Mapping Set Name", "Target 1", "Target 2", "Target 3"],
      [["Balance Sheet Mapping", "BS100", "Current Assets", "Assets"],
       ["Balance Sheet Mapping", "BS200", "Non-Current Assets", "Assets"],
       ["Balance Sheet Mapping", "BS300", "Current Liabilities", "Liabilities"],
       ["P&L Mapping", "PL100", "Revenue", "Income"],
       ["P&L Mapping", "PL200", "Operating Expense", "Expense"]])

sheet("GL Mapping",
      ["Mapping Set Name", "Segment 1", "Segment 2", "Account Description", "Target 1", "Target 2", "Target 3"],
      [["Balance Sheet Mapping", "01", "110000", "Cash at Bank", "BS100", "Current Assets", "Assets"],
       ["Balance Sheet Mapping", "01", "120000", "Trade Receivables", "BS100", "Current Assets", "Assets"],
       ["Balance Sheet Mapping", "01", "150000", "Property Plant & Equipment", "BS200", "Non-Current Assets", "Assets"],
       ["Balance Sheet Mapping", "01", "210000", "Trade Payables", "BS300", "Current Liabilities", "Liabilities"],
       ["P&L Mapping", "01", "400000", "Product Sales", "PL100", "Revenue", "Income"],
       ["P&L Mapping", "01", "610000", "Salaries", "PL200", "Operating Expense", "Expense"]])

wb.save(out)
print(out)
